"""Modules can contribute report sections without editing the report service."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document

from geotek.api import registry
from geotek.core.errors import GeoTekError
from geotek.core.modules import ModuleManifest, ReportEntry
from geotek.services import report


def _manifest(entry_id: str = "demo:one", **overrides) -> ModuleManifest:
    def entries(project):
        return [ReportEntry(entry_id, "Demo section", "ready")]

    def writer(w, project, entry_id):
        w.heading(f"demo written {entry_id}")

    fields = dict(id="demo", version="0.1.0", title_key="demo.title",
                  description_key="demo.description", icon="tool", order=99,
                  report_entries=entries, report_writer=writer)
    fields.update(overrides)
    return ModuleManifest(**fields)


@pytest.fixture
def with_demo(monkeypatch):
    monkeypatch.setattr(registry, "discover", lambda: [_manifest()])


def test_entries_and_writer_must_be_set_together():
    with pytest.raises(GeoTekError):
        _manifest(report_writer=None)


def test_module_section_is_listed(with_demo, monkeypatch):
    project = SimpleNamespace(
        meta=SimpleNamespace(name="P"), soil_profiles=[], structures=[],
        investigations=SimpleNamespace(cpts=[], boreholes=[]))
    monkeypatch.setattr(report, "_liquefaction_status", lambda p: "no_result")
    listed = {s.id: s for s in report.list_sections(project)}
    assert listed["demo:one"].kind == "module"
    assert listed["demo:one"].status == "ready"


def test_entry_id_needs_module_prefix(monkeypatch):
    monkeypatch.setattr(registry, "discover", lambda: [_manifest("other:one")])
    with pytest.raises(GeoTekError):
        report._module_sections(SimpleNamespace())


def test_module_section_is_rendered(with_demo):
    from geotek.core.models import Project, ProjectMeta

    project = Project(meta=ProjectMeta(name="P"))
    data = report.build_docx(project, ["demo:one"])
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "demo written demo:one" in text


def test_pyinstaller_helper_covers_every_module():
    import importlib.util
    from pathlib import Path

    path = Path(__file__).resolve().parents[2] / "tools" / "pyinstaller_modules.py"
    spec = importlib.util.spec_from_file_location("pyinstaller_modules", path)
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    hidden = helper.hidden_imports()
    for id in registry.module_ids():
        assert f"geotek.modules.{id}" in hidden
    assert "geotek.modules._index" in hidden


def _anchored_project(*, analysis=True, anchor=True):
    from geotek.core.models import (
        DesignParameters, GroundwaterCondition, Project, ProjectMeta, SoilLayer, SoilProfile)
    from geotek.core.models.anchors import (
        AnalysisRecord, AnchoredWallAnalysisInputs, AnchoredWallState, AnchorDesignCommon,
        AnchorDesignInputs, AnchorLevel, AnchorRecord, TendonSpec)
    from geotek.core.models.structures import StructureSoilModel
    from geotek.core.module_state import set_state
    from geotek.engine import anchored_wall as engine
    from geotek.engine.anchors import design_anchor, inputs_hash

    project = Project(meta=ProjectMeta(name="Anchors", language="tr"))
    soil = StructureSoilModel(profile=SoilProfile(
        id="SP", gwt=GroundwaterCondition(depth_insitu=100.0), layers=[
            SoilLayer(id="S", top=0, bottom=40, gamma=18.0, gamma_sat=20.0,
                      design=DesignParameters(phi_eff=30.0))]))
    common = AnchorDesignCommon(
        safety_factor=2.0, ultimate_transfer_load=100.0, lock_off_ratio=1.0,
        tendon=TendonSpec(type="strand"))
    state = AnchoredWallState()
    if analysis:
        inputs = AnchoredWallAnalysisInputs(
            excavation_height=10.0, soil=soil, passive_factor=1.5, anchor_design=common,
            levels=[AnchorLevel(depth=d, spacing=2.5, inclination_deg=15.0) for d in (1.5, 4.5, 7.5)])
        state.analysis = AnalysisRecord(inputs=inputs, result=engine.run_analysis(inputs),
                                        result_hash=engine.inputs_hash(inputs))
    if anchor:
        design = AnchorDesignInputs(
            excavation_height=10.0, head_depth=6.25, inclination_deg=15.0, phi_deg=33.0,
            horizontal_load_per_m=172.0, spacing=2.5, **common.model_dump())
        state.anchor = AnchorRecord(inputs=design, result=design_anchor(design),
                                    result_hash=inputs_hash(design))
    set_state(project, "anchored_wall", state, 1)
    return project


def test_anchored_wall_report_entries_follow_the_stored_results():
    from geotek.core.models import Project, ProjectMeta
    from geotek.services import report

    entries = {s.id: s.status for s in report.list_sections(Project(meta=ProjectMeta(name="E")))}
    assert entries["anchored_wall:analysis"] == "no_result"
    assert entries["anchored_wall:anchor"] == "no_result"
    project = _anchored_project()
    entries = {s.id: s.status for s in report.list_sections(project)}
    assert entries["anchored_wall:analysis"] == entries["anchored_wall:anchor"] == "ready"
    project.modules["anchored_wall"].data["analysis"]["inputs"]["excavation_height"] = 11.0
    entries = {s.id: s.status for s in report.list_sections(project)}
    assert entries["anchored_wall:analysis"] == "stale"
    assert entries["anchored_wall:anchor"] == "ready"
    with pytest.raises(report.ReportError):
        report.build_docx(project, ["anchored_wall:analysis"])


@pytest.mark.parametrize("language,heading", [("tr", "Ankrajlı duvar analizi"),
                                              ("en", "Anchored wall analysis")])
def test_anchored_wall_report_is_written_in_both_languages(language, heading):
    from geotek.services import report

    project = _anchored_project()
    data = report.build_docx(project, ["anchored_wall:analysis", "anchored_wall:anchor"],
                             language=language)
    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heading in headings
    cells = {c.text for table in doc.tables for row in table.rows for c in row.cells}
    # three anchor levels with their design loads and tendon counts appear in the levels table
    assert any("×" in text for text in cells)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "anchored_wall." not in text and not any("anchored_wall." in c for c in cells)
    # the envelope type id in assumption texts is translated, never left as a placeholder or raw id
    assert "{kind}" not in text
    expected = "(kum)" if language == "tr" else "(sand)"
    assert f"Görünür basınç zarfı {expected}" in text or f"pressure envelope {expected}" in text


def test_anchored_wall_report_embeds_section_and_diagram_figures():
    from geotek.services import report

    data = report.build_docx(_anchored_project(), ["anchored_wall:analysis"], language="en")
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 2
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Anchored wall section (to scale)" in text and "could not be drawn" not in text
