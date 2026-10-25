"""Report module: section listing, stale refusal and docx content."""

import io

import pytest
from docx import Document

from geotek.core.models.investigations import (
    Borehole, CPTSounding, Investigations, LabRecord, SPTRecord)
from geotek.core.models.profile import GroundwaterCondition
from geotek.core.models.project import ProjectMeta
from geotek.core.models.structures import RetainingWallStructure
from geotek.core.models.walls import ShearKey

from .job_helpers import poll_job, sheet_project, submit, wall_project

BASE = "/api/m/report"
LIQ = {"method": "bi2014", "amax_g": 0.4, "mw": 7.5, "target_fs": 1.1,
       "default_fc": 5, "c_fc": 0, "ic_cutoff": 2.6}
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _text(response) -> str:
    doc = Document(io.BytesIO(response.content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


@pytest.fixture
def open_project(shell, two_layer_profile):
    """Wall + sheet pile + liquefaction results, all current."""
    project = wall_project(geometry=True)
    pile = sheet_project().structures[0]
    pile.id = "sheet"
    project.structures.append(pile)
    project.soil_profiles.append(two_layer_profile)
    project.active_profile_id = two_layer_profile.id
    project.investigations = Investigations(
        boreholes=[Borehole(id="BH1", spt=[SPTRecord(depth=1.5, n_field=6),
                                           SPTRecord(depth=5, n_field=8)],
                            lab=[LabRecord(depth=2, FC=60, LL=40, PI=20)])],
        cpts=[CPTSounding(id="CPT1", depth=[1, 2, 4, 5, 7, 9], qc=[3000, 3000, 3000, 3000, 800, 800],
                          fs=[20, 20, 20, 20, 35, 35])])
    project.meta = ProjectMeta(name="Rapor Projesi", language="en")
    shell.app.state.session.open_project(project)
    for path in ("/api/m/walls/structures/wall/check", "/api/m/sheet_pile/structures/sheet/design"):
        assert poll_job(shell, submit(shell, path))["status"] == "done"
    assert poll_job(shell, submit(shell, "/api/m/liquefaction/run", json=LIQ))["status"] == "done"
    return project


def test_requires_open_project(shell):
    assert shell.get(f"{BASE}/sections").status_code == 409
    assert shell.post(f"{BASE}/docx", json={}).status_code == 409


def test_sections_report_status(shell, open_project):
    sections = {s["id"]: s for s in shell.get(f"{BASE}/sections").json()}
    assert sections["project"]["status"] == "ready"
    assert sections["soil:P1"]["kind"] == "soil_profile"
    assert sections["structure:wall"]["status"] == "ready"
    assert sections["structure:sheet"]["status"] == "ready"
    assert sections["liquefaction"]["status"] == "ready"


def test_default_report_contains_results(shell, open_project):
    response = shell.post(f"{BASE}/docx", json={})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == DOCX
    assert 'filename="Rapor_Projesi_report.docx"' in response.headers["content-disposition"]
    text = _text(response)
    assert "Rapor Projesi" in text and "Checks" in text and "Utilisation" in text
    assert "Max moment" in text and "Liquefaction analysis" in text and "BH1" in text
    wall_checks = open_project.structure("wall").result["checks"]
    assert wall_checks[0]["name"] in text


def test_language_override_and_selection(shell, open_project):
    response = shell.post(f"{BASE}/docx", json={"sections": ["project"], "language": "tr"})
    assert response.status_code == 200
    text = _text(response)
    assert "Proje bilgileri" in text and "Kontroller" not in text


def test_selection_errors(shell, open_project):
    assert shell.post(f"{BASE}/docx", json={"sections": ["structure:nope"]}).status_code == 404
    assert shell.post(f"{BASE}/docx", json={"sections": []}).status_code == 422


def test_stale_wall_is_refused_and_skipped_by_default(shell, open_project):
    wall = open_project.structure("wall")
    assert isinstance(wall, RetainingWallStructure)
    wall.inputs.surcharge += 5
    status = {s["id"]: s["status"] for s in shell.get(f"{BASE}/sections").json()}
    assert status["structure:wall"] == "stale"
    refused = shell.post(f"{BASE}/docx", json={"sections": ["structure:wall"]})
    assert refused.status_code == 409 and refused.json()["code"] == "stale_result"
    assert refused.json()["details"]["sections"] == ["structure:wall"]
    default = _text(shell.post(f"{BASE}/docx", json={}))
    assert "Checks" not in default and "Max moment" in default


def test_missing_result_is_refused(shell, open_project):
    open_project.structure("sheet").result = None
    open_project.structure("sheet").result_hash = None
    refused = shell.post(f"{BASE}/docx", json={"sections": ["structure:sheet"]})
    assert refused.status_code == 409 and refused.json()["code"] == "no_result"


def test_liquefaction_goes_stale_when_investigations_change(shell, open_project):
    open_project.investigations.boreholes[0].spt[0].n_field = 30
    status = {s["id"]: s["status"] for s in shell.get(f"{BASE}/sections").json()}
    assert status["liquefaction"] == "stale"
    refused = shell.post(f"{BASE}/docx", json={"sections": ["liquefaction"]})
    assert refused.status_code == 409 and refused.json()["code"] == "stale_result"


def test_classification_section(shell, open_project):
    response = shell.post(f"{BASE}/docx", json={"sections": ["classification"]})
    assert response.status_code == 200, response.text
    text = _text(response)
    assert "Robertson SBTn" in text and "CPT1" in text and "USCS" in text and "CL" in text


def test_classification_without_data_is_no_result(shell, open_project):
    open_project.investigations.cpts.clear()
    open_project.investigations.boreholes[0].lab.clear()
    refused = shell.post(f"{BASE}/docx", json={"sections": ["classification"]})
    assert refused.status_code == 409 and refused.json()["code"] == "no_result"


def test_document_carries_stored_numeric_values(shell, open_project):
    text = _text(shell.post(f"{BASE}/docx", json={"sections": ["structure:wall", "structure:sheet"]}))
    pile = open_project.structure("sheet").result
    assert f"{pile['M_max']:.1f}" in text and f"{pile['D']:.2f}" in text
    check = open_project.structure("wall").result["checks"][0]
    assert f"{check['value']:.3f}" in text


def test_stale_sheet_pile_is_refused(shell, open_project):
    open_project.structure("sheet").retained_height += 1
    status = {s["id"]: s["status"] for s in shell.get(f"{BASE}/sections").json()}
    assert status["structure:sheet"] == "stale"
    refused = shell.post(f"{BASE}/docx", json={"sections": ["structure:sheet"]})
    assert refused.status_code == 409 and refused.json()["code"] == "stale_result"
    assert refused.json()["details"]["sections"] == ["structure:sheet"]


def test_non_finite_structure_input_is_stale_not_a_server_error(shell, open_project):
    open_project.structure("wall").inputs.surcharge = float("nan")
    sections = shell.get(f"{BASE}/sections")
    assert sections.status_code == 200
    assert {s["id"]: s["status"] for s in sections.json()}["structure:wall"] == "stale"
    refused = shell.post(f"{BASE}/docx", json={"sections": ["structure:wall"]})
    assert refused.status_code == 409 and refused.json()["code"] == "stale_result"
    assert shell.post(f"{BASE}/docx", json={}).status_code == 200  # default selection skips it


def test_non_finite_investigation_value_is_stale_not_a_server_error(shell, open_project):
    open_project.investigations.cpts[0].qc[0] = float("nan")
    sections = shell.get(f"{BASE}/sections")
    assert sections.status_code == 200
    assert {s["id"]: s["status"] for s in sections.json()}["liquefaction"] == "stale"
    refused = shell.post(f"{BASE}/docx", json={"sections": ["liquefaction"]})
    assert refused.status_code == 409 and refused.json()["code"] == "stale_result"


@pytest.mark.parametrize("name, stem", [("Çözüm Şişli", "Cozum_Sisli"), ("", "report"), ("  ", "report")])
def test_file_name_is_ascii_for_turkish_or_empty_project_names(shell, open_project, name, stem):
    open_project.meta.name = name
    response = shell.post(f"{BASE}/docx", json={"sections": ["project"]})
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert disposition.isascii() and disposition.endswith(f'{stem}_report.docx"')


def test_untranslated_warning_codes_are_localised_in_the_document(shell, open_project):
    open_project.meta.language = "tr"
    sections = ["structure:sheet", "liquefaction"]
    tr_text = _text(shell.post(f"{BASE}/docx", json={"sections": sections}))
    assert "liquefaction.warning." not in tr_text and "sheet_pile.warning." not in tr_text
    assert "Tension zone" not in tr_text and "Default FC" not in tr_text
    assert "Laboratuvar FC değeri olmayan" in tr_text
    assert "Default FC = 5.0%" in _text(shell.post(f"{BASE}/docx", json={"sections": ["liquefaction"], "language": "en"}))


def test_docx_shows_names_and_translated_identifiers(shell, open_project):
    """System check 2026-10-20: no raw profile id, enum ids or factor status in the report."""
    open_project.meta.name = "Proje\x01 A"  # control characters cannot go into docx XML
    response = shell.post(f"{BASE}/docx", json={"language": "en"})
    assert response.status_code == 200, response.text
    text = _text(response)
    profile = open_project.profile(open_project.active_profile_id)
    doc = Document(io.BytesIO(response.content))
    rows = {row.cells[0].text: row.cells[-1].text for t in doc.tables for row in t.rows}
    assert rows["Soil profile"] == (profile.name or "(unnamed profile)")
    assert "Boulanger & Idriss (2014)" in text
    for raw in ("\nbi2014\n", "\ncantilever\n", "\nunverified\n", "\nbasic\n", "\nasd\n"):
        assert raw not in f"\n{text}\n"
    assert "Proje A" in text


def test_liquefaction_without_soundings_is_refused(shell, two_layer_profile):
    project = wall_project()
    project.soil_profiles.append(two_layer_profile)
    project.active_profile_id = two_layer_profile.id
    shell.app.state.session.open_project(project)
    response = shell.post("/api/m/liquefaction/run", json=LIQ)
    assert response.status_code == 422
    assert response.json()["details"]["message_code"] == "liquefaction_no_data"


@pytest.mark.parametrize("lang, slope_label, basis, assumption", [
    ("en", "Backfill slope", "Inclined surface", "Surcharge given per inclined surface area"),
    ("tr", "Arka zemin eğimi", "Eğimli yüzey", "Sürşarj eğimli yüzey alanına göre verildi"),
])
def test_slope_and_surcharge_basis_are_reported(shell, open_project, lang, slope_label, basis, assumption):
    open_project.meta.language = lang
    for structure in open_project.structures:
        if structure.kind == "sheet_pile":
            structure.backfill_slope, structure.surcharge, structure.surcharge_basis = 10.0, 15.0, "inclined"
        else:
            structure.backfill_slope = 10.0
            structure.inputs.surcharge, structure.inputs.surcharge_basis = 15.0, "inclined"
    for path in ("/api/m/walls/structures/wall/check", "/api/m/sheet_pile/structures/sheet/design"):
        assert poll_job(shell, submit(shell, path))["status"] == "done"
    text = _text(shell.post(f"{BASE}/docx", json={"sections": ["structure:wall", "structure:sheet"]}))
    assert slope_label in text and basis in text
    assert text.count(assumption) == 2
    assert "report.assumption" not in text and "sheet_pile.assumption" not in text


@pytest.mark.parametrize("lang, labels", [
    ("en", ("Overturning", "Sliding", "Static", "Stem height", "Shear key results", "Under heel")),
    ("tr", ("Devrilme", "Kayma", "Statik", "Gövde yüksekliği", "Kayma dişi sonuçları", "Topuk altı")),
])
def test_wall_report_uses_labels_and_lists_shear_key(shell, open_project, lang, labels):
    wall = open_project.structures[0]
    wall.geometry.key = ShearKey(position="heel", width=0.4, depth=0.5)
    shell.app.state.session.open_project(open_project)
    assert poll_job(shell, submit(shell, "/api/m/walls/structures/wall/check"))["status"] == "done"
    response = shell.post(f"{BASE}/docx", json={"language": lang})
    assert response.status_code == 200, response.text
    text = _text(response)
    for label in labels:
        assert label in text, label
    cells = [c.text for line in Document(io.BytesIO(response.content)).tables for row in line.rows for c in row.cells]
    for raw in ("stem_height", "overturning", "static", "global_stability"):
        assert raw not in cells, raw


@pytest.mark.parametrize("lang, labels", [
    ("en", ("Structure design code (used in this calculation)", "ASD — Global factor of safety",
            "Structure analysis level", "Advanced", "SDS (structure input)",
            "Retained-side groundwater depth, design [m]", "Front-side groundwater depth, seismic [m]",
            "Load per anchor [kN/anchor]", "Anchor spacing [m]", "Anchor inclination from horizontal [°]")),
    ("tr", ("Yapı tasarım kodu (bu hesapta kullanılan)", "ASD — Global güvenlik sayısı",
            "Yapı analiz düzeyi", "İleri", "SDS (yapı girdisi)",
            "Tutulan taraf yeraltı suyu derinliği, tasarım [m]", "Ön taraf yeraltı suyu derinliği, deprem [m]",
            "Ankraj başına yük [kN/ankraj]", "Ankraj aralığı [m]", "Ankraj eğimi (yataydan) [°]")),
])
def test_structure_sections_state_their_own_code_level_water_and_anchor_geometry(
        shell, open_project, lang, labels):
    pile = open_project.structures[1]
    pile.support, pile.anchor_depth = "anchored", 1.0
    pile.anchor_spacing, pile.anchor_inclination = 2.5, 15.0
    pile.analysis_level, pile.sds = "advanced", 0.4
    pile.soil.gwt_retained = GroundwaterCondition(depth_insitu=2, depth_design=1.5, depth_seismic=1.0)
    pile.soil.gwt_front = GroundwaterCondition(depth_insitu=4, depth_design=3.5, depth_seismic=3.0)
    shell.app.state.session.open_project(open_project)
    assert poll_job(shell, submit(shell, "/api/m/sheet_pile/structures/sheet/design"))["status"] == "done"
    response = shell.post(f"{BASE}/docx", json={"language": lang, "sections": ["structure:wall", "structure:sheet"]})
    assert response.status_code == 200, response.text
    text = _text(response)
    for label in labels:
        assert label in text, label
    for value in ("2.50", "15.0", "1.50", "1.00", "3.50", "3.00", "0.400"):
        assert value in text, value
    assert text.count("report.struct") == 0


def test_wall_assumptions_follow_the_report_language(shell, open_project):
    en = _text(shell.post(f"{BASE}/docx", json={"language": "en", "sections": ["structure:wall"]}))
    assert "Concrete unit weight" in en and "Beton birim hacim" not in en
    assert "report.assumption" not in en
    tr = _text(shell.post(f"{BASE}/docx", json={"language": "tr", "sections": ["structure:wall"]}))
    assert "Beton birim hacim ağırlığı" in tr


def test_report_has_cover_summary_and_running_footer(shell, open_project):
    response = shell.post(f"{BASE}/docx", json={"language": "en"})
    doc = Document(io.BytesIO(response.content))
    text = _text(response)
    assert "Geotechnical Calculation Report" in text
    assert "Prepared by" in text and "Checked by" in text
    assert "Utilisation is the load (demand)" in text  # summary explains the number it shows
    assert "All design checks are satisfied" in text
    assert "PAGE" in doc.sections[0].footer._element.xml


def test_report_embeds_wall_and_pile_figures(shell, open_project):
    response = shell.post(f"{BASE}/docx", json={"language": "tr"})
    doc = Document(io.BytesIO(response.content))
    assert len(doc.inline_shapes) >= 5  # wall section, pressures, check bars, pile diagrams
    text = _text(response)
    assert "Şekil 1." in text and "Eğilme momenti" not in text.split("Şekil 1.")[0]


def test_figures_return_none_without_data():
    from geotek.services import report_figures as figures

    assert figures.sheet_pile_diagrams({"z": [], "net": [], "shear": [], "moment": []}, {}) is None
    assert figures.pressure_plot([], "static", {}) is None
    assert figures.check_bars([], {}) is None
    assert figures.wall_section("cantilever", {}, slope_deg=0, surcharge=0, gwt_depth=None, labels={}) is None
