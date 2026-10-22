from common import *
from docx import Document
import zipfile,re,hashlib,csv,collections
meta={};docs={}
for lang in ['tr','en']:
 f=OUT/'docx'/f'qa-all-{lang}.docx';d=Document(f);docs[lang]=d
 txt='\n'.join([p.text for p in d.paragraphs]+[f'TABLE {i}\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows) for i,t in enumerate(d.tables)])
 (OUT/'docx'/f'qa-all-{lang}.txt').write_text(txt)
 with zipfile.ZipFile(f) as z:
  xml=z.read('word/document.xml');(OUT/'docx'/f'qa-all-{lang}.xml').write_bytes(xml)
 meta[lang]={'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'valid_zip':zipfile.is_zipfile(f),'xml_bytes':len(xml),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'table_shapes':[[len(t.rows),len(t.columns)] for t in d.tables],'headings':[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]}
 print(lang,meta[lang]);print(txt[-6000:])
 with (OUT/'docx'/f'qa-all-{lang}-tables.csv').open('w') as f:
  w=csv.writer(f);w.writerow(['table','row','cells']);
  for i,t in enumerate(d.tables):
   for j,row in enumerate(t.rows):w.writerow([i,j,json.dumps([c.text for c in row.cells],ensure_ascii=False)])
# Compare table cell numbers only, including translated descriptive cells for transparency.
num=lambda s:re.findall(r'(?<![A-Za-z])[-+]?\d+(?:[.,]\d+)?',s)
diffs=[];equal=0
for i,(tr,en) in enumerate(zip(docs['tr'].tables,docs['en'].tables)):
 for j,(rtr,ren) in enumerate(zip(tr.rows,en.rows)):
  for k,(ctr,cen) in enumerate(zip(rtr.cells,ren.cells)):
   a,b=num(ctr.text),num(cen.text)
   if a==b:equal+=1
   else:diffs.append({'table':i,'row':j,'cell':k,'tr':ctr.text,'en':cen.text,'numbers_tr':a,'numbers_en':b})
meta['comparison']={'table_shapes_equal':meta['tr']['table_shapes']==meta['en']['table_shapes'],'numeric_cells_equal':equal,'numeric_differences':diffs}
(OUT/'docx-inspection.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2));print('COMPARE',json.dumps(meta['comparison'],ensure_ascii=False)[:7000])
