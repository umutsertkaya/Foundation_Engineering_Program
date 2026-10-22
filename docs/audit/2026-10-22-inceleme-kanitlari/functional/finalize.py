from common import *
from docx import Document
import zipfile,csv,hashlib,collections
for jid in ['0cb285e1-18a8-40d8-a95a-fb6a8ba07804','39588d72-132f-44bd-91dc-5c348ae1666d']:api('GET','/api/jobs/'+jid)
api('POST','/api/project/save')
selected=['report.md','scenarios.csv','requests.csv','jobs-summary.json','docx-inspection.json','docx-dto-comparison.json','wall-draft-loss.json','discard-job.json','cancel-job.json','pile-cancel-job.json','narrow-metrics.json','envelope-seismic.json','envelope-static.json','docx/qa-all-tr.docx','docx/qa-all-en.docx','docx/qa-all-en-tables.csv','docx/qa-all-tr-tables.csv']
selected += ['screenshots/'+x+'.png' for x in ['04-wall-wizard-geometry','08-wall-missing-factor-block','09-wall-advanced-check','18-pile-False-result','18-pile-True-result','20-anchor-multilevel-result','22-anchor-design-result','24-liquefaction-Boulanger-result','28-report-en-downloaded','29-wall-en-dark-narrow-viewport','32-liq-cpt-en-dark-narrow-viewport','36-pile-saved-input-stale','37-envelope-seismic-inputs-ui','38-envelope-seismic-result-ui','39-envelope-static-result-ui','44-report-stale-disabled','46-wall-job-discarded-after-reopen','47-wall-cancelled-real-job','49-wall-global-save-local-draft','50-wall-global-save-reload-lost']]
missing=[s for s in selected if not (OUT/s).is_file()];assert not missing,missing
(OUT/'selected-evidence.json').write_text(json.dumps([{'path':s,'bytes':(OUT/s).stat().st_size,'sha256':hashlib.sha256((OUT/s).read_bytes()).hexdigest()}for s in selected],ensure_ascii=False,indent=2))
for lang in ['tr','en']:
 f=OUT/'docx'/f'qa-all-{lang}.docx';assert zipfile.is_zipfile(f);d=Document(f);assert len(d.tables)==25 and len(d.paragraphs)==163
assert json.loads((OUT/'docx-dto-comparison.json').read_text())['all_pass']
assert json.loads((OUT/'discard-job.json').read_text())['error']['details']['discarded'] is True
assert json.loads((OUT/'cancel-job.json').read_text())['job']['status']=='cancelled'
print('Selected evidence files:',len(selected),'screenshots total:',len(list((OUT/'screenshots').glob('*.png'))),'report bytes:',(OUT/'report.md').stat().st_size)
