from common import *
from docx import Document
import csv,re,collections,math,hashlib
# Keep complete raw responses; provide compact searchable request, job and output inventories.
rows=[];jobs={}
for fn in ['api.jsonl','browser-network.jsonl']:
 for ln,line in enumerate((OUT/fn).read_text().splitlines(),1):
  r=json.loads(line);body=r.get('response');url=r.get('url','');method=r.get('method');path=url.replace(BASE,'')
  d=body if isinstance(body,dict) else {}
  rows.append({'file':fn,'line':ln,'method':method,'url':path,'http_status':r.get('status'),'code':d.get('code',''),'job_id':d.get('job_id',''),'job_status':d.get('status','') if 'job_id' in d else '', 'discarded':d.get('error',{}).get('details',{}).get('discarded','') if isinstance(d.get('error'),dict) else ''})
  if d.get('job_id') and d.get('status') in ['done','failed','cancelled']:
   res=d.get('result')or{};j={k:d.get(k) for k in ['job_id','kind','status','error']};j.update(result_keys=list(res),D=res.get('D'),M_max=res.get('M_max'),result_hash=res.get('result_hash'))
   if 'per_sounding'in res:j['soundings']={k:{'n':len(v['depth']),'minFS':min(x for x in v['FS'] if x is not None),'n_liquefiable':v['status'].count('liquefiable')}for k,v in res['per_sounding'].items()}
   jobs[d['job_id']]=j
with (OUT/'requests.csv').open('w')as f:
 w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
(OUT/'jobs-summary.json').write_text(json.dumps(list(jobs.values()),ensure_ascii=False,indent=2))
# Verify numeric result delivery to DOCX at documented display rounding.
doc=Document(OUT/'docx/qa-all-en.docx');checks=[]
def check(ti,label,expected,tol):
 row=next(r for r in doc.tables[ti].rows if r.cells[0].text==label);actual=row.cells[1].text
 tol=0.5 * 10**(-len(actual.split('.')[1])) + 1e-10 if '.' in actual else tol
 try:ok=abs(float(actual)-expected)<=tol
 except:ok=False
 checks.append({'table':ti,'label':label,'DTO':expected,'DOCX':actual,'tolerance':tol,'pass':ok})
for filename,ti in [('pile-console-url.txt',11),('pile-anchored-url.txt',12)]:
 sid=(OUT/filename).read_text().split('/s/')[1];d=client.get('/api/structures/'+sid).json()['result']
 for label,key,tol in [('Theoretical embedment D0 [m]','D0',.00051),('Design embedment D [m]','D',.00051),('Total length [m]','total_length',.00051),('Max moment [kNm/m]','M_max',.00051)]:check(ti,label,d[key],tol)
state=client.get('/api/m/anchored_wall/state').json()['state'];a=state['anchor']['result'];aw=state['analysis']['result']
for label,key,tol in [('Design load T (kN)','design_load',.051),('Bond length (m)','bond_length',.0051),('Unbonded length (m)','unbonded_length',.0051),('Total length (m)','total_length',.0051)]:
 if key in a:check(23,label,a[key],tol)
for filename,ti in [('wall-url.txt',9),('gravity-url.txt',15)]:
 sid=(OUT/filename).read_text().split('/s/')[1];d=client.get('/api/structures/'+sid).json()['result']
 for row,item in zip(doc.tables[ti].rows[1:],d['checks']):
  val=row.cells[2].text;tol=.00051
  checks.append({'table':ti,'label':row.cells[0].text+' '+row.cells[1].text,'DTO':item['value'],'DOCX':val,'tolerance':tol,'pass':abs(float(val)-item['value'])<=tol})
for label,key in [('Theoretical embedment D0 (m)','embedment_theoretical'),('Design embedment D (m)','embedment'),('Wall length (m)','wall_length'),('Base reaction R_B (kN/m)','base_reaction'),('Max moment (kNm/m)','m_max'),('Max shear (kN/m)','v_max')]:check(20,label,aw[key],.0051)
for row,level in zip(doc.tables[22].rows[1:],aw['levels']):
 for idx,key in [(3,'horizontal_load_per_m'),(4,'axial_load_per_anchor')]:
  val=row.cells[idx].text
  checks.append({'table':22,'label':key+' level '+row.cells[0].text,'DTO':level[key],'DOCX':val,'tolerance':.051,'pass':abs(float(val)-level[key])<=.051})
# Explicit unit tokens in report; no conversion inferred from string alone.
texts='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
unit_tokens=['[m]','[kPa]','[kN/m³]','[kNm/m]','[kN/m]','[cm³/m]','(kN)','(m)','(°)']
(OUT/'docx-dto-comparison.json').write_text(json.dumps({'checks':checks,'all_pass':all(c['pass']for c in checks),'unit_tokens':{x:x in texts for x in unit_tokens}},ensure_ascii=False,indent=2));print('requests',len(rows),'terminal jobs',len(jobs),'docx DTO numeric checks',len(checks),'allpass',all(c['pass']for c in checks))
manifest=[{'file':str(x.relative_to(OUT)),'bytes':x.stat().st_size,'sha256':hashlib.sha256(x.read_bytes()).hexdigest()}for x in sorted(OUT.rglob('*'))if x.is_file()and not any(v in x.parts for v in ['browser','__pycache__'])]
(OUT/'artifact-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
